"""
طبقة الوصول الوحيدة لملف قاعدة البيانات Excel (subscribers_1.xlsx).

كل قراءة أو كتابة في الملف تمر عبر هذه الطبقة فقط، حتى يسهل لاحقاً
الانتقال إلى قاعدة بيانات حقيقية دون تغيير بقية الكود.

الضمانات:
- أسماء الأوراق والأعمدة محفوظة كما هي، ولا تُحذف أي صفوف.
- تُتجاهل أي ورقة أخرى في الملف ولا تُقرأ ولا تُعدَّل.
- كل عملية كتابة تتم داخل قفل (قفل ملف بين العمليات + قفل خيط داخل العملية)
  ثم حفظ آمن: ملف مؤقت ثم استبدال ذري — لمنع تلف الملف.
- نسخة احتياطية يومية تُؤخذ تلقائياً قبل أول كتابة في اليوم.
"""
from __future__ import annotations

import os
import secrets
import shutil
import threading
import time
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import openpyxl

from common import DT_FMT, excel_safe, now_str, parse_dt, excel_safe as _safe

try:  # ويندوز
    import msvcrt
except ImportError:  # لينكس/ماك
    msvcrt = None
    import fcntl

# ---------------------------------------------------------------- الأوراق
USERS = "Users"
CODES = "code_1"
LOGS = "Logs"
EXTENSIONS = "Extensions"
RESETS = "Resets"

# الأوراق المسموح بها فقط — أي ورقة أخرى تُتجاهل تماماً
SHEET_HEADERS = {
    USERS: [
        "ID", "username", "email", "phone", "password_hash", "salt", "role",
        "status", "created_at", "activated_at", "expires_at", "last_login",
        "disabled_at", "notes",
    ],
    CODES: ["code", "status", "used_by_phone", "used_at"],
    LOGS: ["التاريخ", "الفاعل", "الإجراء", "الهدف", "التفاصيل", "النتيجة"],
    EXTENSIONS: [
        "التاريخ", "الأدمن", "رقم الجوال المشترك", "تاريخ الانتهاء قبل",
        "تاريخ الانتهاء بعد", "عدد الأيام", "السبب",
    ],
    RESETS: ["phone", "code_hash", "expires_at", "attempts"],
}

# تحقق من أسماء الأعمدة الفعلية في ملف المستخدم (يبقى الملف هو المرجع)
EXT_HEADERS_ON_DISK = [
    "التاريخ", "الأدمن", "رقم جوال المشترك", "تاريخ الانتهاء قبل",
    "تاريخ الانتهاء بعد", "عدد الأيام", "السبب",
]
SHEET_HEADERS[EXTENSIONS] = EXT_HEADERS_ON_DISK

_LOCK_TIMEOUT = 20.0  # أقصى انتظار لقفل الملف بالثواني


def _check_sheet(sheet: str) -> None:
    if sheet not in SHEET_HEADERS:
        raise KeyError(f"ورقة غير مسموح بها: {sheet}")


# ---------------------------------------------------------------- قفل الملف
class _FileLock:
    """قفل ملف بين العمليات، مع إعادة الدخول الآمن للخيط نفسه."""

    def __init__(self, path: Path, timeout: float = _LOCK_TIMEOUT):
        self.path = Path(path)
        self.timeout = timeout
        self._fh = None
        self._owner = None
        self._depth = 0
        self._mutex = threading.Lock()

    def acquire(self) -> None:
        with self._mutex:
            ident = threading.get_ident()
            if self._owner == ident:  # إعادة دخول من نفس الخيط
                self._depth += 1
                return
        # انتظار القفل خارج mutex حتى لا نحجب خيوطاً أخرى
        deadline = time.time() + self.timeout
        fh = None
        while True:
            try:
                fh = open(self.path, "a+b")
                _os_lock(fh)
                break
            except OSError:
                if fh is not None:
                    try:
                        fh.close()
                    except OSError:
                        pass
                    fh = None
                if time.time() > deadline:
                    raise TimeoutError(
                        "تعذّر الوصول لقاعدة البيانات: الملف مقفول من عملية أخرى. حاول بعد لحظات."
                    )
                time.sleep(0.05)
        with self._mutex:
            self._fh = fh
            self._owner = threading.get_ident()
            self._depth = 1

    def release(self) -> None:
        with self._mutex:
            if self._owner != threading.get_ident():
                return
            self._depth -= 1
            if self._depth > 0:
                return
            fh, self._fh = self._fh, None
            self._owner = None
        if fh is not None:
            _os_unlock(fh)
            try:
                fh.close()
            except OSError:
                pass


def _os_lock(fh) -> None:
    if msvcrt is not None:
        fh.seek(0, os.SEEK_END)
        if fh.tell() == 0:
            fh.write(b"\0")
            fh.flush()
        fh.seek(0)
        msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
    else:
        fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)


def _os_unlock(fh) -> None:
    try:
        if msvcrt is not None:
            fh.seek(0)
            msvcrt.locking(fh.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            fcntl.flock(fh.fileno(), fcntl.LOCK_UN)
    except OSError:
        pass


# قفل خيط واحد مشترك لعمليات الكتابة داخل نفس العملية
_PROCESS_LOCK = threading.RLock()


# ---------------------------------------------------------------- مساعدات الأوراق
def norm(value) -> str:
    """توحيد قيمة مقروءة من Excel إلى نص."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(DT_FMT)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def sheet_rows(wb, sheet: str) -> list[dict]:
    """يقرأ صفوف ورقة كقائمة قواميس (يتخطى الصفوف الفارغة تماماً)."""
    _check_sheet(sheet)
    ws = wb[sheet]
    header = [norm(c.value) for c in ws[1]]
    rows: list[dict] = []
    for idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        row = {"_row": idx}
        for key, value in zip(header, values):
            if key:
                row[key] = norm(value)
        rows.append(row)
    return rows


def _header_map(ws) -> list[str]:
    return [norm(c.value) for c in ws[1]]


def _last_used_row(ws) -> int:
    """رقم آخر صف يحمل أي بيانات (تجنّب الصفوف الفارغة المنسيّة في نهاية الورقة)."""
    for idx in range(ws.max_row, 1, -1):
        for cell in ws[idx]:
            if cell.value not in (None, ""):
                return idx
    return 1


def append_row(wb, sheet: str, data: dict) -> int:
    """يضيف صفوفاً جديداً بقيم محيّدة (بدون حذف أي صف قائم). يعيد رقم الصف."""
    _check_sheet(sheet)
    ws = wb[sheet]
    row_num = _last_used_row(ws) + 1
    for col, key in enumerate(_header_map(ws), start=1):
        if key and key in data:
            value = data[key]
            if isinstance(value, str):
                value = _safe(value)
            ws.cell(row=row_num, column=col, value=value)
    return row_num


def set_row(wb, sheet: str, row_number: int, data: dict) -> None:
    """يحدّث خلايا صف موجود (لا يحذف ولا يُنشئ صفوفاً)."""
    _check_sheet(sheet)
    ws = wb[sheet]
    for col, key in enumerate(_header_map(ws), start=1):
        if key and key in data:
            value = data[key]
            if isinstance(value, str):
                value = _safe(value)
            ws.cell(row=row_number, column=col, value=value)


def log_row(wb, actor, action, target, details="", result="") -> None:
    """يسجّل حدثاً في ورقة Logs داخل معاملة مفتوحة (يجب استدعاؤها داخل write)."""
    append_row(wb, LOGS, {
        "التاريخ": now_str(),
        "الفاعل": str(actor or ""),
        "الإجراء": str(action or ""),
        "الهدف": str(target or ""),
        "التفاصيل": str(details or ""),
        "النتيجة": str(result or ""),
    })


# ---------------------------------------------------------------- الطبقة الرئيسية
class ExcelStore:
    """الواجهة الوحيدة لقراءة/كتابة ملف المشتركين."""

    def __init__(self, data_file: str | Path, backup_dir: str | Path):
        self.data_file = Path(data_file)
        self.backup_dir = Path(backup_dir)
        self.lock_file = self.data_file.with_name(self.data_file.name + ".lock")
        self._thread_local = threading.local()
        self._file_lock = _FileLock(self.lock_file)

    # ---------- إنشاء الملف والأوراق
    def init_file(self) -> None:
        """ينشئ الملف والأوراق بالمخطط الصحيح إن لم توجد."""
        self.data_file.parent.mkdir(parents=True, exist_ok=True)
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        if not self.data_file.exists():
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            for name, headers in SHEET_HEADERS.items():
                ws = wb.create_sheet(name)
                ws.append(headers)
            self._save_atomic(wb)
            return
        with self.read() as wb:
            missing = [n for n in SHEET_HEADERS if n not in wb.sheetnames]
        if missing:
            with self.write() as wb:
                for name in missing:
                    ws = wb.create_sheet(name)
                    ws.append(SHEET_HEADERS[name])

    # ---------- عمليات القراءة والكتابة
    @contextmanager
    def read(self):
        """يقرأ الملف داخل قفل (بدون كتابة)."""
        self._file_lock.acquire()
        _PROCESS_LOCK.acquire()
        try:
            wb = openpyxl.load_workbook(self.data_file)
            try:
                yield wb
            finally:
                wb.close()
        finally:
            _PROCESS_LOCK.release()
            self._file_lock.release()

    @contextmanager
    def write(self):
        """
        يفتح معاملة كتابة داخل قفل: تُطبَّق التعديلات على wb ثم يُحفظ الملف
        آمناً (ملف مؤقت + استبدال ذري). أي استثناء يعني إلغاء التعديلات.
        ممنوع فتح معاملة كتابة متداخلة — استخدم append_row/set_row/log_row على wb.
        """
        if getattr(self._thread_local, "writing", False):
            raise RuntimeError(
                "محاولة فتح معاملة كتابة متداخلة داخل معاملة أخرى — "
                "استخدم append_row/set_row/log_row على wb مباشرة."
            )
        self._thread_local.writing = True
        self._file_lock.acquire()
        _PROCESS_LOCK.acquire()
        wb = None
        try:
            # نسخة احتياطية يومية قبل أول تعديل في اليوم
            self.ensure_backup()
            wb = openpyxl.load_workbook(self.data_file)
            ok = False
            try:
                yield wb
                ok = True
            finally:
                if ok:
                    self._save_atomic(wb)
                else:
                    wb.close()
        finally:
            if wb is not None and not getattr(self._thread_local, "writing", True):
                pass
            _PROCESS_LOCK.release()
            self._file_lock.release()
            self._thread_local.writing = False

    def _save_atomic(self, wb) -> None:
        """حفظ آمن: ملف مؤقت في نفس المجلد ثم استبدال ذري."""
        tmp = self.data_file.with_name(f".{self.data_file.name}.{secrets.token_hex(6)}.tmp")
        try:
            wb.save(tmp)
        finally:
            try:
                wb.close()  # إغلاق أرشيف الملف الأصلي (ضروري على ويندوز)
            except Exception:
                pass
        last_error = None
        for _ in range(5):  # إعادة محاولة قليلة إذا كان الملف مؤقتاً مستخدماً
            try:
                os.replace(tmp, self.data_file)
                last_error = None
                break
            except OSError as exc:  # PermissionError على ويندوز مثلاً
                last_error = exc
                time.sleep(0.1)
        if last_error is not None:
            if tmp.exists():
                tmp.unlink(missing_ok=True)
            raise last_error

    # ---------- النسخ الاحتياطي اليومي
    def ensure_backup(self) -> Path | None:
        """نسخة احتياطية واحدة يومياً (تُستدعى قبل أول كتابة في اليوم)."""
        if not self.data_file.exists():
            return None
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y-%m-%d")
        target = self.backup_dir / f"{self.data_file.stem}_{stamp}{self.data_file.suffix}"
        if not target.exists():
            shutil.copy2(self.data_file, target)
        self._prune_backups(keep=30)
        return target

    def _prune_backups(self, keep: int = 30) -> None:
        try:
            files = sorted(
                (p for p in self.backup_dir.glob(f"{self.data_file.stem}_*{self.data_file.suffix}")),
                key=lambda p: p.name,
            )
            for old in files[:-keep] if len(files) > keep else []:
                old.unlink(missing_ok=True)
        except OSError:
            pass

    # ---------- واجهات مختصرة للطبقات العليا
    def rows(self, sheet: str) -> list[dict]:
        with self.read() as wb:
            return sheet_rows(wb, sheet)

    def append(self, sheet: str, data: dict) -> None:
        with self.write() as wb:
            append_row(wb, sheet, data)

    def update(self, sheet: str, row_number: int, data: dict) -> None:
        with self.write() as wb:
            set_row(wb, sheet, row_number, data)

    def log(self, actor, action, target, details="", result="") -> None:
        """يسجّل حدثاً في Logs (عملية مستقلة بقفل وحفظ خاص بها)."""
        with self.write() as wb:
            log_row(wb, actor, action, target, details, result)

    # ---------- استعلامات شائعة
    def all_users(self) -> list[dict]:
        """كل حسابات حقيقية (لها ID) — لا يحذف أي صف ولا يُنشئ غيره."""
        return [r for r in self.rows(USERS) if r.get("ID")]

    def find_user(self, **conditions) -> dict | None:
        """يبحث عن حساب بمطابقة دقيقة (غير حساسة لحالة الأحرف)."""
        target = {k: str(v).strip().lower() for k, v in conditions.items()}
        for row in self.all_users():
            if all(str(row.get(k, "")).strip().lower() == v for k, v in target.items()):
                return row
        return None

    def find_code(self, code: str) -> dict | None:
        """يبحث عن كود في ورقة code_1 (مطابقة غير حساسة لحالة الأحرف)."""
        wanted = str(code or "").strip().lower()
        for row in self.rows(CODES):
            if row.get("code", "").strip().lower() == wanted:
                return row
        return None

    def all_codes(self) -> list[dict]:
        return [r for r in self.rows(CODES) if r.get("code")]

    def sweep_expired(self) -> int:
        """
        يحوّل كل حساب «مفعل» قد انتهى تاريخه إلى «منتهي».
        يُستدعى عند تسجيل الدخول وعبر المهمة اليومية المجدولة.
        """
        with self.read() as wb:
            rows = sheet_rows(wb, USERS)
        targets = [
            r["_row"] for r in rows
            if r.get("ID") and r.get("status") == "مفعل"
            and (dt := parse_dt(r.get("expires_at"))) and datetime.now() >= dt
        ]
        if not targets:
            return 0
        changed = 0
        with self.write() as wb:
            for row_num in targets:
                set_row(wb, USERS, row_num, {"status": "منتهي"})
                changed += 1
            log_row(wb, "النظام", "فحص الانتهاء اليومي", "-", f"تحويل {changed} حساب إلى منتهي", "تم")
        return changed
