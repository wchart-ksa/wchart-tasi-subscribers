"""تقرير التفعيل: ملف Excel (رقم الجوال + حالة التفعيل) يتحدّث مع تغيّر الحالة."""
import os

import openpyxl

import app as app_module
import support
from common import ROLE_ADMIN, STATUS_ACTIVE, STATUS_DISABLED, STATUS_INACTIVE
from storage import activation_report
from storage.excel_store import CODES


def test_rows_filter_and_status_mapping(store):
    support.seed_user(store, "act1", "0501111111", "Secret123", status=STATUS_ACTIVE)
    support.seed_user(store, "dis1", "0502222222", "Secret123", status=STATUS_DISABLED)
    support.seed_user(store, "inact1", "0503333333", "Secret123",
                      status=STATUS_INACTIVE, activated_at="")

    assert dict(activation_report.rows_for(store)) == {
        "0501111111": "مفعل",
        "0502222222": "غير مفعل",
    }


def test_sync_writes_readable_two_column_file(store, tmp_path):
    support.seed_user(store, "act2", "0502222222", "Secret123", status=STATUS_ACTIVE)
    path = tmp_path / "rep.xlsx"
    activation_report.sync(store, path)

    assert path.exists()
    ws = openpyxl.load_workbook(path)[activation_report.SHEET]
    values = list(ws.iter_rows(values_only=True))
    assert values[0] == ("رقم الجوال", "حالة التفعيل")
    assert ("0502222222", "مفعل") in values[1:]
    assert ws.max_column == 2


def test_report_created_on_activation_and_updated_on_toggle(client, store, tmp_path, monkeypatch):
    report_path = tmp_path / "rep.xlsx"
    monkeypatch.setitem(app_module.CFG, "ACTIVATION_REPORT_FILE", str(report_path))

    support.seed_user(store, "web1", "0504444444", "Secret123",
                      status=STATUS_INACTIVE, activated_at="")
    store.append(CODES, {"code": "WEB1234", "status": "", "used_by_phone": "", "used_at": ""})
    support.login(client, "web1", "Secret123")
    support.post(client, "/user/activate", {"code": "WEB1234"}, follow=True)

    created = list(openpyxl.load_workbook(report_path)[activation_report.SHEET].iter_rows(values_only=True))
    assert ("0504444444", "مفعل") in created

    support.seed_user(store, "adm1", "0509990000", "AdminPass1", role=ROLE_ADMIN)
    support.login(client, "adm1", "AdminPass1")
    support.post(client, "/admin/users/toggle/0504444444")

    updated = list(openpyxl.load_workbook(report_path)[activation_report.SHEET].iter_rows(values_only=True))
    assert ("0504444444", "غير مفعل") in updated
    assert ("0504444444", "مفعل") not in updated


def test_admin_can_download_report(client, store, tmp_path, monkeypatch):
    monkeypatch.setitem(app_module.CFG, "ACTIVATION_REPORT_FILE", str(tmp_path / "rep.xlsx"))
    support.seed_user(store, "adm3", "0505555555", "AdminPass1", role=ROLE_ADMIN)
    support.login(client, "adm3", "AdminPass1")
    resp = client.get("/admin/activation-report")
    assert resp.status_code == 200
    assert "spreadsheet" in resp.headers["Content-Type"]


def test_report_auto_refreshes_when_data_changes(client, store, tmp_path, monkeypatch):
    report = tmp_path / "auto.xlsx"
    monkeypatch.setitem(app_module.CFG, "ACTIVATION_REPORT_FILE", str(report))

    client.get("/auth/login")  # before_request ينشئ التقرير تلقائياً
    assert report.exists()

    support.seed_user(store, "late1", "0508888888", "Secret123", status=STATUS_ACTIVE)
    os.utime(report, (0, 0))  # اجعل التقرير أقدم من ملف البيانات

    client.get("/auth/login")
    values = list(openpyxl.load_workbook(report)[activation_report.SHEET].iter_rows(values_only=True))
    assert ("0508888888", "مفعل") in values

